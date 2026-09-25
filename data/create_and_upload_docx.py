import requests
import json
import os
import sys
import base64
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

def get_auth_header():
    config_path = os.path.expanduser(r"~\.gemini\config\mcp_config.json")
    with open(config_path, "r") as f:
        config = json.load(f)
    return config["mcpServers"]["mnema"]["headers"]["Authorization"]

def markdown_to_docx(md_path, docx_path):
    doc = Document()
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue
        if line_str.startswith("# "):
            doc.add_heading(line_str[2:], level=1)
        elif line_str.startswith("## "):
            doc.add_heading(line_str[3:], level=2)
        elif line_str.startswith("### "):
            doc.add_heading(line_str[4:], level=3)
        elif line_str.startswith("- ") or line_str.startswith("* "):
            doc.add_paragraph(line_str[2:], style='List Bullet')
        else:
            doc.add_paragraph(line_str)

    doc.save(docx_path)

DOCS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "mnema_docs")

docs_to_upload = [
    {
        "md_path": os.path.join(DOCS_DIR, "Gen_UI_System_Architecture_and_Overview.md"),
        "folder_id": "b025205e-7c9f-4806-a44f-7c507960d563",
        "docx_filename": "Gen_UI_System_Architecture_and_Overview.docx"
    },
    {
        "md_path": os.path.join(DOCS_DIR, "Gen_UI_Generative_UI_Engine_and_Component_Registry.md"),
        "folder_id": "b025205e-7c9f-4806-a44f-7c507960d563",
        "docx_filename": "Gen_UI_Generative_UI_Engine_and_Component_Registry.docx"
    },
    {
        "md_path": os.path.join(DOCS_DIR, "Gen_UI_API_Key_Vault_and_Proxy_Security.md"),
        "folder_id": "b025205e-7c9f-4806-a44f-7c507960d563",
        "docx_filename": "Gen_UI_API_Key_Vault_and_Proxy_Security.docx"
    },
    {
        "md_path": os.path.join(DOCS_DIR, "Gen_UI_Database_Schema_and_Data_Models.md"),
        "folder_id": "b025205e-7c9f-4806-a44f-7c507960d563",
        "docx_filename": "Gen_UI_Database_Schema_and_Data_Models.docx"
    },
    {
        "md_path": os.path.join(DOCS_DIR, "Gen_UI_Architectural_Decisions_Log.md"),
        "folder_id": "d78b1bdb-5d76-4a92-a753-6a7acf608c8d",
        "docx_filename": "Gen_UI_Architectural_Decisions_Log.docx"
    },
    {
        "md_path": os.path.join(DOCS_DIR, "Gen_UI_Master_Build_Spec_and_Deployment.md"),
        "folder_id": "79db8f07-d035-4105-9310-64e22307fd20",
        "docx_filename": "Gen_UI_Master_Build_Spec_and_Deployment.docx"
    }
]

def main():
    url = "https://api.theboringpeople.in/mcp"
    try:
        auth_header = get_auth_header()
    except Exception as e:
        print(f"Failed to read mcp_config.json: {e}")
        return

    headers = {
        "Authorization": auth_header,
        "Accept": "application/json, text/event-stream",
        "Content-Type": "application/json"
    }

    for idx, item in enumerate(docs_to_upload):
        md_path = item["md_path"]
        folder_id = item["folder_id"]
        docx_filename = item["docx_filename"]
        temp_docx_path = os.path.join(os.path.dirname(md_path), docx_filename)

        if not os.path.exists(md_path):
            print(f"File not found: {md_path}")
            continue

        print(f"Converting {os.path.basename(md_path)} to docx...")
        markdown_to_docx(md_path, temp_docx_path)

        with open(temp_docx_path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode("utf-8")

        payload = {
            "jsonrpc": "2.0",
            "id": idx + 1,
            "method": "tools/call",
            "params": {
                "name": "upload_doc_file",
                "arguments": {
                    "filename": docx_filename,
                    "content_base64": b64,
                    "folder_id": folder_id
                }
            }
        }

        print(f"Uploading {docx_filename} to folder {folder_id}...")
        r = requests.post(url, headers=headers, json=payload)
        print(f"Status: {r.status_code}")
        print(r.text)
        print("-" * 50)

        # Cleanup temp docx
        if os.path.exists(temp_docx_path):
            os.remove(temp_docx_path)

if __name__ == "__main__":
    main()
